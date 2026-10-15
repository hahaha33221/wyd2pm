"""Q&A DB 엑셀(WYD_2nd_Prep_Meeting_QA_DB.xlsx)을 앱 데이터(src/constants/seedData.js)와
VPS 데이터베이스용 SQL(db/seed.sql)로 바꿉니다.

사용:
    python3 scripts/import_db.py "WYD_2nd_Prep_Meeting_QA_DB.xlsx"

시트 3개를 읽습니다 (왼쪽 네비게이션도 이 시트 구성을 따릅니다).
  - QA_DB        : No / 질문번호 / 대분류코드 / 대분류 / 중분류번호 / 중분류 / 소분류코드 / 소분류 / 질문 / 답변
  - Opinions     : No / 관련 항목 / 의견
  - Meeting_Info : 항목 / 내용
엑셀이 원본입니다. 내용을 고칠 때는 엑셀을 고친 뒤 이 스크립트를 다시 실행하세요.

다른 언어는 translations/<언어코드>/*.json 에서 읽어 합칩니다 (한국어는 엑셀 그대로, 바꾸지 않음).
  - qa         : {"Q1": {"q": "...", "a": "..."}}
  - sections   : {"Ⅰ": "..."}            topics: {"Ⅰ-1": "..."}   subtopics: {"Ⅰ-1-a": "..."}
  - meeting    : {"회의명": {"label": "...", "value": "..."}}
  - opinions   : {"1": "..."}            opinionTopics: {"숙소": "..."}
필요: pip install openpyxl
"""
import json
import re
import sys
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parent.parent
OUT = ROOT / "src" / "constants" / "seedData.js"
TRANSLATIONS = ROOT / "translations"
SQL_OUT = ROOT / "db" / "seed.sql"


def clean(value, one_line=False):
    """보이지 않는 공백을 없애고 앞뒤 공백을 정리합니다."""
    if value is None:
        return ""
    s = str(value).replace("\xa0", " ").replace("​", "").replace("﻿", "")
    s = s.replace("\r\n", "\n").replace("\r", "\n")
    lines = [re.sub(r"[ \t]+", " ", line).strip() for line in s.split("\n")]
    s = "\n".join(lines).strip()
    if one_line:
        s = re.sub(r"\s*\n\s*(?=·)", "", s)
        s = re.sub(r"\s*\n\s*", " ", s)
    return s


def rows(ws, header):
    """머리글을 확인하고 내용 행을 dict로 돌려줍니다."""
    it = ws.iter_rows(values_only=True)
    first = [clean(c) for c in next(it, [])]
    if first[: len(header)] != header:
        sys.exit(f"'{ws.title}' 시트 머리글이 다릅니다: {first} (기대: {header})")
    for row_no, row in enumerate(it, start=2):
        vals = [clean(v) for v in (list(row) + [None] * len(header))[: len(header)]]
        if any(vals):
            yield row_no, dict(zip(header, vals))


def q(v):
    """SQL 문자열 값 (NULL 포함)."""
    if v is None or v == "":
        return "NULL"
    return "'" + str(v).replace("'", "''") + "'"


def build_sql(qa, opinions, info, source):
    """엑셀 내용과 DB를 똑같이 맞추는 SQL. 엑셀에 없는 행은 지웁니다 (참가자 질문 접수 inquiry는 건드리지 않음)."""
    out = [
        "-- 자동 생성 파일입니다. 직접 고치지 말고 엑셀을 고친 뒤 다시 만드세요.",
        f"-- 원본: {source}",
        "-- 적용: psql -v ON_ERROR_STOP=1 -f db/schema.sql -f db/seed.sql",
        "BEGIN;",
        "",
        "-- 회의 정보",
    ]
    for i, r in enumerate(info):
        out.append(
            f"INSERT INTO meeting_info (label, value, sort_order) VALUES ({q(r['label']['ko'])}, {q(r['value']['ko'])}, {i}) "
            "ON CONFLICT (label) DO UPDATE SET value = EXCLUDED.value, sort_order = EXCLUDED.sort_order;"
        )
    labels = ", ".join(q(r["label"]["ko"]) for r in info) or "NULL"
    out.append(f"DELETE FROM meeting_info WHERE label NOT IN ({labels});")

    out += ["", "-- 분류: 대분류 > 중분류 > 소분류"]
    seen_s, seen_t, seen_st = {}, {}, {}
    for it in qa:
        s_code, s_name = it["sectionNo"], it["section"]["ko"]
        t_no, t_name = it["topicNo"], it["topic"]["ko"]
        if s_code not in seen_s:
            seen_s[s_code] = len(seen_s)
            out.append(
                f"INSERT INTO qa_section (code, name, sort_order) VALUES ({q(s_code)}, {q(s_name)}, {seen_s[s_code]}) "
                "ON CONFLICT (code) DO UPDATE SET name = EXCLUDED.name, sort_order = EXCLUDED.sort_order;"
            )
        tk = (s_code, t_no)
        if tk not in seen_t:
            seen_t[tk] = len(seen_t)
            out.append(
                "INSERT INTO qa_topic (section_id, no, name, sort_order) "
                f"SELECT id, {q(t_no)}, {q(t_name)}, {seen_t[tk]} FROM qa_section WHERE code = {q(s_code)} "
                "ON CONFLICT (section_id, no) DO UPDATE SET name = EXCLUDED.name, sort_order = EXCLUDED.sort_order;"
            )
        if it.get("sub"):
            sk = (s_code, t_no, it["subNo"])
            if sk not in seen_st:
                seen_st[sk] = len(seen_st)
                out.append(
                    "INSERT INTO qa_subtopic (topic_id, code, name, sort_order) "
                    f"SELECT t.id, {q(it['subNo'])}, {q(it['sub']['ko'])}, {seen_st[sk]} FROM qa_topic t "
                    f"JOIN qa_section s ON s.id = t.section_id WHERE s.code = {q(s_code)} AND t.no = {q(t_no)} "
                    "ON CONFLICT (topic_id, code) DO UPDATE SET name = EXCLUDED.name, sort_order = EXCLUDED.sort_order;"
                )

    out += ["", "-- 질문·답변"]
    for i, it in enumerate(qa):
        s_code, t_no = it["sectionNo"], it["topicNo"]
        sub_id = (
            f"(SELECT st.id FROM qa_subtopic st WHERE st.topic_id = t.id AND st.code = {q(it['subNo'])})"
            if it.get("sub")
            else "NULL"
        )
        out.append(
            "INSERT INTO qa_item (question_no, section_id, topic_id, subtopic_id, sort_order) "
            f"SELECT {q(it['no'])}, s.id, t.id, {sub_id}, {i} FROM qa_section s "
            f"JOIN qa_topic t ON t.section_id = s.id AND t.no = {q(t_no)} WHERE s.code = {q(s_code)} "
            "ON CONFLICT (question_no) DO UPDATE SET section_id = EXCLUDED.section_id, topic_id = EXCLUDED.topic_id, "
            "subtopic_id = EXCLUDED.subtopic_id, sort_order = EXCLUDED.sort_order, updated_at = now();"
        )
        for lang, question in it["q"].items():
            out.append(
                "INSERT INTO qa_item_i18n (item_id, lang, question, answer) "
                f"SELECT id, {q(lang)}, {q(question)}, {q(it['a'].get(lang))} FROM qa_item WHERE question_no = {q(it['no'])} "
                "ON CONFLICT (item_id, lang) DO UPDATE SET question = EXCLUDED.question, answer = EXCLUDED.answer, updated_at = now();"
            )
        langs = ", ".join(q(lang) for lang in it["q"])
        out.append(
            f"DELETE FROM qa_item_i18n WHERE item_id = (SELECT id FROM qa_item WHERE question_no = {q(it['no'])}) AND lang NOT IN ({langs});"
        )
    nos = ", ".join(q(it["no"]) for it in qa) or "NULL"
    out += [
        f"DELETE FROM qa_item WHERE question_no NOT IN ({nos});",
        "DELETE FROM qa_subtopic st WHERE NOT EXISTS (SELECT 1 FROM qa_item i WHERE i.subtopic_id = st.id);",
        "DELETE FROM qa_topic t WHERE NOT EXISTS (SELECT 1 FROM qa_item i WHERE i.topic_id = t.id);",
        "DELETE FROM qa_section s WHERE NOT EXISTS (SELECT 1 FROM qa_item i WHERE i.section_id = s.id);",
        "",
        "-- 의견",
    ]
    for i, o in enumerate(opinions):
        out.append(
            f"INSERT INTO opinion (no, topic, content, sort_order) VALUES ({o['no']}, {q(o['topic']['ko'])}, {q(o['text']['ko'])}, {i}) "
            "ON CONFLICT (no) DO UPDATE SET topic = EXCLUDED.topic, content = EXCLUDED.content, sort_order = EXCLUDED.sort_order;"
        )
    op_nos = ", ".join(str(o["no"]) for o in opinions) or "NULL"
    out += [f"DELETE FROM opinion WHERE no NOT IN ({op_nos});", "", "COMMIT;", ""]
    return "\n".join(out)


def load_translations():
    """translations/<lang>/*.json 을 언어별로 합칩니다."""
    out = {}
    if not TRANSLATIONS.is_dir():
        return out
    for lang_dir in sorted(p for p in TRANSLATIONS.iterdir() if p.is_dir()):
        merged = {}
        for f in sorted(lang_dir.glob("*.json")):
            for key, val in json.loads(f.read_text(encoding="utf-8")).items():
                merged.setdefault(key, {}).update(val)
        out[lang_dir.name] = merged
    return out


def apply_translations(qa, opinions, info, tr, problems):
    """한국어(ko)는 그대로 두고 다른 언어 값만 더합니다."""
    for lang, t in tr.items():
        if lang == "ko":
            continue
        missing = 0
        for it in qa:
            tk = f"{it['sectionNo']}-{it['topicNo']}"
            if it["sectionNo"] in t.get("sections", {}):
                it["section"][lang] = t["sections"][it["sectionNo"]]
            if tk in t.get("topics", {}):
                it["topic"][lang] = t["topics"][tk]
            if it.get("sub") and f"{tk}-{it['subNo']}" in t.get("subtopics", {}):
                it["sub"][lang] = t["subtopics"][f"{tk}-{it['subNo']}"]
            row = t.get("qa", {}).get(it["no"])
            if row and row.get("q"):
                it["q"][lang] = row["q"]
                if row.get("a") and it["a"]:
                    it["a"][lang] = row["a"]
            else:
                missing += 1
        for o in opinions:
            if o["topic"]["ko"] in t.get("opinionTopics", {}):
                o["topic"][lang] = t["opinionTopics"][o["topic"]["ko"]]
            if str(o["no"]) in t.get("opinions", {}):
                o["text"][lang] = t["opinions"][str(o["no"])]
        for r in info:
            m = t.get("meeting", {}).get(r["label"]["ko"])
            if m:
                r["label"][lang] = m.get("label", "")
                r["value"][lang] = m.get("value", "")
        if missing:
            problems.append(f"{lang}: 번역이 없는 질문 {missing}개 (화면에는 한국어로 표시)")


def sheet(wb, name):
    if name not in wb.sheetnames:
        sys.exit(f"'{name}' 시트가 없습니다. 시트 목록: {wb.sheetnames}")
    return wb[name]


def main():
    if len(sys.argv) != 2:
        sys.exit(__doc__)
    wb = openpyxl.load_workbook(sys.argv[1], data_only=True)
    problems = []

    qa = []
    qa_header = ["No", "질문번호", "대분류코드", "대분류", "중분류번호", "중분류", "소분류코드", "소분류", "질문", "답변"]
    for row_no, r in rows(sheet(wb, "QA_DB"), qa_header):
        if not (r["질문번호"] and r["대분류"] and r["중분류"] and r["질문"]):
            problems.append(f"QA_DB {row_no}행: 질문번호·대분류·중분류·질문 중 빈 칸이 있어 건너뜀")
            continue
        if not r["답변"]:
            problems.append(f"QA_DB {row_no}행: 답변이 비어 있음 (화면에는 '답변 준비 중')")
        item = {
            "id": r["질문번호"],
            "no": r["질문번호"],
            "section": {"ko": r["대분류"]},
            "sectionNo": r["대분류코드"],
            "topic": {"ko": r["중분류"]},
            "topicNo": r["중분류번호"],
            # 질문 앞의 '○ ' 같은 표시는 뺍니다
            "q": {"ko": re.sub(r"^○\s*", "", r["질문"])},
            "a": {"ko": r["답변"]} if r["답변"] else {},
        }
        if r["소분류"]:
            item["sub"] = {"ko": r["소분류"]}
            item["subNo"] = r["소분류코드"]
        qa.append(item)

    opinions = [
        {
            "id": f"O{r['No'] or row_no}",
            "no": int(r["No"]) if r["No"].isdigit() else row_no,
            "topic": {"ko": clean(r["관련 항목"], one_line=True)},
            "text": {"ko": r["의견"]},
        }
        for row_no, r in rows(sheet(wb, "Opinions"), ["No", "관련 항목", "의견"])
        if r["의견"]
    ]

    info = [
        {"label": {"ko": r["항목"]}, "value": {"ko": r["내용"]}}
        for _, r in rows(sheet(wb, "Meeting_Info"), ["항목", "내용"])
        if r["항목"]
    ]

    apply_translations(qa, opinions, info, load_translations(), problems)

    def js(v):
        return json.dumps(v, ensure_ascii=False, indent=2)

    OUT.write_text(
        "// 자동 생성 파일입니다. 직접 고치지 말고 엑셀을 고친 뒤 다시 만드세요.\n"
        f"// 원본: {Path(sys.argv[1]).name}\n"
        "// 만드는 법: python3 scripts/import_db.py \"<엑셀 파일>\"\n\n"
        "// QA_DB 시트\n"
        f"export const SEED = {js(qa)};\n\n"
        "// Opinions 시트\n"
        f"export const OPINIONS = {js(opinions)};\n\n"
        "// Meeting_Info 시트\n"
        f"export const MEETING_INFO = {js(info)};\n",
        encoding="utf-8",
    )

    SQL_OUT.parent.mkdir(exist_ok=True)
    SQL_OUT.write_text(build_sql(qa, opinions, info, Path(sys.argv[1]).name), encoding="utf-8")

    print(f"질문 {len(qa)}개 · 의견 {len(opinions)}개 · 회의 정보 {len(info)}개 → {OUT.relative_to(ROOT)}, {SQL_OUT.relative_to(ROOT)}")
    for p in problems:
        print("  ! " + p)


if __name__ == "__main__":
    main()
