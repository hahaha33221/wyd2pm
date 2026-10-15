"""엑셀(세션별 질문 및 답변)을 앱의 FAQ 데이터(src/constants/seedData.js)로 바꿉니다.

사용:
    python3 scripts/import_xlsx.py "2PM_세션별 질문 및 답변_KOR_v0.2.xlsx"

엑셀의 '질문' 시트 열: 구분 / 주제 / 세부 항목 / 질문 / 답변
엑셀이 원본입니다. 답변을 고칠 때는 엑셀을 고친 뒤 이 스크립트를 다시 실행하세요.
필요: pip install openpyxl
"""
import json
import re
import sys
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parent.parent
OUT = ROOT / "src" / "constants" / "seedData.js"
HEADER = ["구분", "주제", "세부 항목", "질문", "답변"]


def clean(value):
    """보이지 않는 공백을 없애고 앞뒤 공백을 정리합니다."""
    if value is None:
        return ""
    s = str(value).replace("\xa0", " ").replace("​", "").replace("﻿", "")
    s = s.replace("\r\n", "\n").replace("\r", "\n")
    lines = [re.sub(r"[ \t]+", " ", line).strip() for line in s.split("\n")]
    return "\n".join(lines).strip()


def split_number(label):
    """'Ⅰ. 참가 신청 및 등록' → ('Ⅰ', '참가 신청 및 등록'), '2.순례자앱' → ('2', '순례자앱')."""
    m = re.match(r"^\s*([ⅠⅡⅢⅣⅤⅥⅦⅧⅨⅩ]+|\d+|[a-z])\s*\.\s*(.*)$", label)
    return (m.group(1), m.group(2).strip()) if m else ("", label.strip())


def find_sheet(wb):
    for ws in wb.worksheets:
        first = [clean(c) for c in next(ws.iter_rows(max_row=1, values_only=True), [])]
        if first[: len(HEADER)] == HEADER:
            return ws
    sys.exit(f"'{' / '.join(HEADER)}' 머리글이 있는 시트를 찾지 못했습니다.")


def main():
    if len(sys.argv) != 2:
        sys.exit(__doc__)
    wb = openpyxl.load_workbook(sys.argv[1], data_only=True)
    ws = find_sheet(wb)

    items, problems = [], []
    for row_no, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        section, topic, sub, q, a = (clean(v) for v in (list(row) + [None] * 5)[:5])
        if not any((section, topic, sub, q, a)):
            continue
        if not (section and topic and q):
            problems.append(f"{row_no}행: 구분·주제·질문 중 빈 칸이 있어 건너뜀")
            continue
        if not a:
            problems.append(f"{row_no}행: 답변이 비어 있음 (화면에는 '답변 준비 중'으로 표시)")

        s_no, s_name = split_number(section)
        t_no, t_name = split_number(topic)
        _, sub_name = split_number(sub) if sub else ("", "")
        item = {
            "id": f"r{row_no}",
            "section": {"ko": s_name},
            "sectionNo": s_no,
            "topic": {"ko": t_name},
            "topicNo": t_no,
            "q": {"ko": q},
            "a": {"ko": a} if a else {},
        }
        if sub_name:
            item["sub"] = {"ko": sub_name}
        items.append(item)

    body = json.dumps(items, ensure_ascii=False, indent=2)
    OUT.write_text(
        "// 자동 생성 파일입니다. 직접 고치지 말고 엑셀을 고친 뒤 다시 만드세요.\n"
        f"// 원본: {Path(sys.argv[1]).name}\n"
        "// 만드는 법: python3 scripts/import_xlsx.py \"<엑셀 파일>\"\n"
        f"export const SEED = {body};\n",
        encoding="utf-8",
    )

    sections = {}
    for it in items:
        sections.setdefault(it["section"]["ko"], set()).add(it["topic"]["ko"])
    print(f"질문 {len(items)}개 → {OUT.relative_to(ROOT)}")
    for name, topics in sections.items():
        n = sum(1 for it in items if it["section"]["ko"] == name)
        print(f"  {name} ({n}개): {', '.join(sorted(topics))}")
    for p in problems:
        print("  ! " + p)


if __name__ == "__main__":
    main()
